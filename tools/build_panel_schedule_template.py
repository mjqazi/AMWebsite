#!/usr/bin/env python3
"""Build downloads/panel-schedule-template.xlsx (run from anywhere).

Sheets: "How to use", "NEC Panel Schedule" (US panelboard, VA per phase),
"AS-NZS DB Schedule" (Australian/NZ distribution board, amps per phase),
hidden "Lists" for drop-downs. Example rows use illustrative values only.
"""
import os
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.formatting.rule import CellIsRule, FormulaRule
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.utils import get_column_letter

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
OUT = os.path.join(ROOT, 'downloads', 'panel-schedule-template.xlsx')
URL = 'advancedmechanix.com/panel-schedule-template.html'

NAVY = '1B3A6B'
thin = Side(style='thin', color='B8B2A4')
BORDER = Border(left=thin, right=thin, top=thin, bottom=thin)
HEAD_FILL = PatternFill('solid', fgColor=NAVY)
LABEL_FILL = PatternFill('solid', fgColor='EFEBDF')
INPUT_FILL = PatternFill('solid', fgColor='FFFFFF')
EX_FILL = PatternFill('solid', fgColor='FFF4D6')
TOT_FILL = PatternFill('solid', fgColor='E4DED1')
WHITE_B = Font(bold=True, color='FFFFFF')
BOLD = Font(bold=True)
WRAP = Alignment(wrap_text=True, vertical='top')
CENTER = Alignment(horizontal='center', vertical='center', wrap_text=True)

wb = Workbook()

# ---------------------------------------------------------------- Lists
lists = wb.create_sheet('Lists')
LISTS = {
    'A': ('Load type (NEC)', ['Lighting', 'Receptacles', 'Motor', 'HVAC', 'Kitchen equipment', 'Continuous', 'Non-continuous', 'Spare', 'Space']),
    'B': ('Poles', [1, 2, 3]),
    'C': ('Phase', ['A', 'B', 'C']),
    'D': ('Device (AS/NZS)', ['MCB', 'RCBO', 'MCCB', 'Fuse', 'Isolator', 'Spare', 'Space']),
    'E': ('Curve', ['B', 'C', 'D', 'n/a']),
    'F': ('Yes/No', ['Y', 'N']),
    'G': ('Phase (AS/NZS)', ['L1', 'L2', 'L3', 'L1-L2-L3']),
    'H': ('Status', ['Proposed', 'IFC', 'Installed', 'As-built']),
}
for col, (title, vals) in LISTS.items():
    lists[col + '1'] = title
    lists[col + '1'].font = BOLD
    for i, v in enumerate(vals, start=2):
        lists[f'{col}{i}'] = v
lists.sheet_state = 'hidden'

def lst(col):
    n = len(LISTS[col][1])
    return f'Lists!${col}$2:${col}${n + 1}'

def dv(ws, col, ref, first, last):
    v = DataValidation(type='list', formula1=lst(col), allow_blank=True)
    v.add(f'{ref}{first}:{ref}{last}')
    ws.add_data_validation(v)

def header_block(ws, title, fields_left, fields_right, ncols):
    ws['A1'] = title
    ws['A1'].font = Font(bold=True, size=14, color=NAVY)
    ws['A2'] = f'Template by Advanced Mechanix — {URL}. Free to use and adapt.'
    ws['A2'].font = Font(italic=True, size=9, color='6A6D75')
    r = 3
    for i in range(max(len(fields_left), len(fields_right))):
        for (lab, col_l, col_v, span) in ((fields_left[i] if i < len(fields_left) else None, 1, 2, 3),
                                          (fields_right[i] if i < len(fields_right) else None, 6, 7, 3)):
            if not lab:
                continue
            c = ws.cell(row=r, column=col_l, value=lab)
            c.font = BOLD; c.fill = LABEL_FILL; c.border = BORDER
            ws.merge_cells(start_row=r, start_column=col_v, end_row=r, end_column=col_v + span - 1)
            for k in range(col_v, col_v + span):
                ws.cell(row=r, column=k).border = BORDER
                ws.cell(row=r, column=k).fill = INPUT_FILL
        r += 1
    return r

def table_head(ws, row, heads, widths):
    for i, (h, w) in enumerate(zip(heads, widths), start=1):
        c = ws.cell(row=row, column=i, value=h)
        c.font = WHITE_B; c.fill = HEAD_FILL; c.alignment = CENTER; c.border = BORDER
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.row_dimensions[row].height = 42

# ---------------------------------------------------------------- NEC sheet
nec = wb.active
nec.title = 'NEC Panel Schedule'
r = header_block(nec, 'PANEL SCHEDULE — PANELBOARD (NEC PROJECTS)',
                 ['Project', 'Client', 'Drawing / document no.', 'Revision', 'Date'],
                 ['Panel name', 'Location', 'Fed from', 'Prepared / checked by', 'Status'], 12)
# system data block
sys_rows = [('System voltage L-L (V)', 208), ('System voltage L-N (V)', 120), ('Phases / wires', '3Ø 4W'),
            ('Bus rating (A)', 225), ('Main: MCB rating (A) or MLO', 'MLO'), ('Short-circuit rating (kA)', None),
            ('Mounting / enclosure', None), ('Imbalance flag limit (%) — set by you', 15)]
sr = r + 1
nec.cell(row=r, column=1, value='System data (EXAMPLE values — replace)').font = Font(bold=True, color=NAVY)
for i, (lab, val) in enumerate(sys_rows):
    c = nec.cell(row=sr + i, column=1, value=lab); c.font = BOLD; c.fill = LABEL_FILL; c.border = BORDER
    nec.merge_cells(start_row=sr + i, start_column=2, end_row=sr + i, end_column=3)
    v = nec.cell(row=sr + i, column=2, value=val); v.border = BORDER; v.fill = EX_FILL if val is not None else INPUT_FILL
    nec.cell(row=sr + i, column=3).border = BORDER
VLL, VLN, LIMIT = f'$B${sr}', f'$B${sr + 1}', f'$B${sr + 7}'
r = sr + len(sys_rows) + 1
NEC_HEAD = ['Ckt no.', 'Description', 'Load type', 'Poles', 'Breaker trip (A)', 'Conductor size (AWG/kcmil)',
            'Phase (pos.)', 'Load on A (VA)', 'Load on B (VA)', 'Load on C (VA)', 'Demand factor (%)', 'Demand load (VA)', 'Status', 'Notes']
NEC_W = [8, 30, 16, 7, 11, 14, 9, 12, 12, 12, 11, 13, 11, 34]
th = r
table_head(nec, th, NEC_HEAD, NEC_W)
first = th + 1
NCKT = 42
last = first + NCKT - 1
for i in range(NCKT):
    row = first + i
    ckt = i + 1
    phase = 'ABC'[(i // 2) % 3]           # standard 2-column panelboard: 1-2 A, 3-4 B, 5-6 C ...
    vals = [ckt, None, None, None, None, None, phase, None, None, None, None,
            f'=IF(COUNT(H{row}:J{row})=0,"",SUM(H{row}:J{row})*IF(ISNUMBER(K{row}),K{row},100)/100)', None, None]
    for j, v in enumerate(vals, start=1):
        c = nec.cell(row=row, column=j, value=v); c.border = BORDER
        if j in (1, 7): c.alignment = CENTER
example = {2: 'EXAMPLE — Lighting, level 1 (illustrative)', 3: 'Lighting', 4: 1, 5: 20, 6: '12', 8: 1440, 11: 100, 13: 'Proposed',
           14: 'EXAMPLE ROW — illustrative values; delete or overwrite'}
for col, v in example.items():
    nec.cell(row=first, column=col, value=v)
for col in range(1, 15):
    nec.cell(row=first, column=col).fill = EX_FILL
dv(nec, 'A', 'C', first, last); dv(nec, 'B', 'D', first, last); dv(nec, 'C', 'G', first, last); dv(nec, 'H', 'M', first, last)
# totals
t = last + 2
tot = [('Connected load per phase (VA)', [f'=SUM(H{first}:H{last})', f'=SUM(I{first}:I{last})', f'=SUM(J{first}:J{last})']),
       ('Current per phase (A) = VA / V L-N', [f'=IF(ISNUMBER({VLN}),ROUND(H{t}/{VLN},1),"")', f'=IF(ISNUMBER({VLN}),ROUND(I{t}/{VLN},1),"")', f'=IF(ISNUMBER({VLN}),ROUND(J{t}/{VLN},1),"")'])]
for k, (lab, fs) in enumerate(tot):
    rr = t + k
    c = nec.cell(row=rr, column=2, value=lab); c.font = BOLD
    for j, f in enumerate(fs):
        cc = nec.cell(row=rr, column=8 + j, value=f); cc.fill = TOT_FILL; cc.border = BORDER; cc.font = BOLD
summ = [
    ('Total connected load (VA)', f'=SUM(H{t}:J{t})'),
    ('Total demand load (VA)', f'=SUM(L{first}:L{last})'),
    ('Demand current, 3-phase (A) = VA / (√3 × V L-L)', f'=IF(ISNUMBER({VLL}),ROUND(B{t + 4}/(SQRT(3)*{VLL}),1),"")'),
    ('Phase imbalance (%) = max deviation from average / average', f'=IF(AVERAGE(H{t}:J{t})>0,ROUND(MAX(ABS(H{t}-AVERAGE(H{t}:J{t})),ABS(I{t}-AVERAGE(H{t}:J{t})),ABS(J{t}-AVERAGE(H{t}:J{t})))/AVERAGE(H{t}:J{t})*100,1),"")'),
    ('Imbalance check', f'=IF(ISNUMBER(B{t + 6}),IF(B{t + 6}<={LIMIT},"OK","CHECK"),"")'),
]
for k, (lab, f) in enumerate(summ):
    rr = t + 3 + k
    a = nec.cell(row=rr, column=1, value=lab); a.font = BOLD; a.fill = LABEL_FILL; a.alignment = WRAP
    nec.row_dimensions[rr].height = 30
    b = nec.cell(row=rr, column=2, value=f); b.fill = TOT_FILL; b.border = BORDER; b.font = BOLD
nec.column_dimensions['A'].width = 30
nec.conditional_formatting.add(f'B{t + 7}', CellIsRule(operator='equal', formula=['"CHECK"'], fill=PatternFill('solid', fgColor='F8D7D0')))
nec.cell(row=t + 9, column=1, value='Load calculations, demand factors and conductor/overcurrent sizing must be verified by a competent person to the NEC edition adopted by the authority having jurisdiction. This template does not reproduce NEC tables.').alignment = WRAP
nec.merge_cells(start_row=t + 9, start_column=1, end_row=t + 9, end_column=10)
nec.row_dimensions[t + 9].height = 32
nec.freeze_panes = f'C{first}'
nec.print_title_rows = f'{th}:{th}'
nec.page_setup.orientation = 'landscape'; nec.page_setup.fitToWidth = 1; nec.page_setup.fitToHeight = 0
nec.sheet_properties.pageSetUpPr.fitToPage = True

# ---------------------------------------------------------------- AS/NZS sheet
au = wb.create_sheet('AS-NZS DB Schedule')
r = header_block(au, 'DISTRIBUTION BOARD SCHEDULE (AS/NZS PROJECTS)',
                 ['Project', 'Client', 'Drawing / document no.', 'Revision', 'Date'],
                 ['Board name', 'Location', 'Supplied from', 'Prepared / checked by', 'Status'], 13)
sys_rows = [('Supply (V)', '400/230 V 3-phase 4-wire 50 Hz'), ('Nominal voltage L-N (V)', 230), ('Main switch rating (A)', 160),
            ('Prospective fault current at board (kA)', None), ('Assembly: form of separation / IP rating', None),
            ('Earthing system', 'MEN'), ('Imbalance flag limit (%) — set by you', 15)]
sr = r + 1
au.cell(row=r, column=1, value='Board data (EXAMPLE values — replace)').font = Font(bold=True, color=NAVY)
for i, (lab, val) in enumerate(sys_rows):
    c = au.cell(row=sr + i, column=1, value=lab); c.font = BOLD; c.fill = LABEL_FILL; c.border = BORDER
    au.merge_cells(start_row=sr + i, start_column=2, end_row=sr + i, end_column=3)
    v = au.cell(row=sr + i, column=2, value=val); v.border = BORDER; v.fill = EX_FILL if val is not None else INPUT_FILL
    au.cell(row=sr + i, column=3).border = BORDER
LIMIT_AU = f'$B${sr + 6}'
r = sr + len(sys_rows) + 1
AU_HEAD = ['Circuit no.', 'Description', 'Device', 'Rating (A)', 'Curve', 'Poles', 'RCD 30 mA (Y/N)', 'Phase',
           'Cable size (mm²)', 'Cable ID (cable schedule)', 'Load L1 (A)', 'Load L2 (A)', 'Load L3 (A)', 'Status', 'Notes']
AU_W = [30, 30, 10, 9, 7, 7, 10, 10, 10, 14, 10, 10, 10, 11, 34]
th = r
table_head(au, th, AU_HEAD, AU_W)
first = th + 1
NC = 36
last = first + NC - 1
for i in range(NC):
    row = first + i
    au.cell(row=row, column=1, value=i + 1).alignment = CENTER
    for j in range(1, 16):
        au.cell(row=row, column=j).border = BORDER
ex = {2: 'EXAMPLE — GPOs, office (illustrative)', 3: 'RCBO', 4: 20, 5: 'C', 6: 1, 7: 'Y', 8: 'L1', 9: 2.5, 10: 'EXAMPLE-001',
      11: 8, 14: 'Proposed', 15: 'EXAMPLE ROW — illustrative values; delete or overwrite'}
for col, v in ex.items():
    au.cell(row=first, column=col, value=v)
for col in range(1, 16):
    au.cell(row=first, column=col).fill = EX_FILL
dv(au, 'D', 'C', first, last); dv(au, 'E', 'E', first, last); dv(au, 'B', 'F', first, last)
dv(au, 'F', 'G', first, last); dv(au, 'G', 'H', first, last); dv(au, 'H', 'N', first, last)
t = last + 2
c = au.cell(row=t, column=2, value='Connected load per phase (A)'); c.font = BOLD
for j, col in enumerate('KLM'):
    cc = au.cell(row=t, column=11 + j, value=f'=SUM({col}{first}:{col}{last})'); cc.fill = TOT_FILL; cc.border = BORDER; cc.font = BOLD
summ = [
    ('Highest phase load (A)', f'=MAX(K{t}:M{t})'),
    ('Phase imbalance (%) = max deviation from average / average', f'=IF(AVERAGE(K{t}:M{t})>0,ROUND(MAX(ABS(K{t}-AVERAGE(K{t}:M{t})),ABS(L{t}-AVERAGE(K{t}:M{t})),ABS(M{t}-AVERAGE(K{t}:M{t})))/AVERAGE(K{t}:M{t})*100,1),"")'),
    ('Imbalance check', f'=IF(ISNUMBER(B{t + 3}),IF(B{t + 3}<={LIMIT_AU},"OK","CHECK"),"")'),
    ('Maximum demand (A) — from your AS/NZS 3000 assessment', None),
    ('Main switch check: max demand ≤ main switch rating', f'=IF(AND(ISNUMBER(B{t + 5}),ISNUMBER($B${sr + 2})),IF(B{t + 5}<=$B${sr + 2},"OK","CHECK"),"")'),
]
for k, (lab, f) in enumerate(summ):
    rr = t + 2 + k
    a = au.cell(row=rr, column=1, value=lab); a.font = BOLD; a.fill = LABEL_FILL; a.alignment = WRAP
    au.row_dimensions[rr].height = 30
    b = au.cell(row=rr, column=2, value=f); b.border = BORDER; b.font = BOLD
    b.fill = INPUT_FILL if f is None else TOT_FILL
for ref in (f'B{t + 4}', f'B{t + 6}'):
    au.conditional_formatting.add(ref, CellIsRule(operator='equal', formula=['"CHECK"'], fill=PatternFill('solid', fgColor='F8D7D0')))
au.cell(row=t + 8, column=1, value='Maximum demand, protective device selection, RCD requirements and cable sizing must be verified by a competent person to AS/NZS 3000, AS/NZS 3008.1.1 and the assembly manufacturer\'s AS/NZS 61439 data. This template does not reproduce tables from any standard.').alignment = WRAP
au.merge_cells(start_row=t + 8, start_column=1, end_row=t + 8, end_column=10)
au.row_dimensions[t + 8].height = 32
au.freeze_panes = f'C{first}'
au.print_title_rows = f'{th}:{th}'
au.page_setup.orientation = 'landscape'; au.page_setup.fitToWidth = 1; au.page_setup.fitToHeight = 0
au.sheet_properties.pageSetUpPr.fitToPage = True

# ---------------------------------------------------------------- How to use
how = wb.create_sheet('How to use', 0)
how.column_dimensions['A'].width = 110
lines = [
    ('How to use this panel schedule template', 'title'),
    (f'Advanced Mechanix — {URL}. Free to use and adapt; provided as-is, not engineering advice.', 'note'),
    ('', None),
    ('Two versions', 'h'),
    ('• "NEC Panel Schedule" — a US panelboard schedule: 42 circuit positions, phase position A/B/C in the usual 1-2 A, 3-4 B, 5-6 C pattern, loads entered in VA per phase, optional demand factor per circuit.', None),
    ('• "AS-NZS DB Schedule" — an Australian / New Zealand distribution board schedule: circuit, protective device (MCB, RCBO, MCCB), rating, curve, RCD, phase L1/L2/L3, cable size in mm² and cable ID, loads entered in amps per phase.', None),
    ('', None),
    ('Steps', 'h'),
    ('1. Fill in the header and the system / board data. Yellow cells are EXAMPLE values — replace them.', None),
    ('2. Enter one row per circuit. Delete or overwrite the yellow EXAMPLE row. Multi-pole circuits: enter the load on each phase they occupy.', None),
    ('3. Most columns have drop-down lists (load type, poles, device, curve, RCD, phase, status).', None),
    ('4. Totals calculate automatically: load per phase, current per phase (NEC sheet), total connected and demand load, and phase imbalance.', None),
    ('5. Set your own imbalance limit in the data block. The check cell shows OK or CHECK against that limit; it is a project choice, not a code value.', None),
    ('6. AS/NZS sheet: enter the maximum demand from your AS/NZS 3000 assessment; the sheet checks it against the main switch rating.', None),
    ('7. Use the Cable ID column to tie each circuit to the cable schedule (free template: advancedmechanix.com/cable-schedule-template.html).', None),
    ('', None),
    ('Standards', 'h'),
    ('US: panelboards and load calculations are governed by the NEC edition your authority having jurisdiction has adopted (load calculations in Article 220, panelboards in Article 408). AU/NZ: wiring rules AS/NZS 3000, cable selection AS/NZS 3008.1.1, switchboard assemblies the AS/NZS 61439 series.', None),
    ('This workbook does not reproduce any table from these standards. Look up demand factors, ratings and cable capacities in your own licensed copy or design software, and have the schedule checked by a competent person.', None),
]
for i, (txt, kind) in enumerate(lines, start=1):
    c = how.cell(row=i, column=1, value=txt)
    c.alignment = Alignment(wrap_text=True, vertical='top')
    if kind == 'title': c.font = Font(bold=True, size=14, color=NAVY)
    elif kind == 'h': c.font = Font(bold=True, color=NAVY)
    elif kind == 'note': c.font = Font(italic=True, size=9, color='6A6D75')

wb.move_sheet('Lists', offset=len(wb.sheetnames))
wb.active = 0
wb.properties.title = 'Panel schedule template'
wb.properties.creator = 'Advanced Mechanix'
os.makedirs(os.path.dirname(OUT), exist_ok=True)
wb.save(OUT)
print('wrote', OUT)
